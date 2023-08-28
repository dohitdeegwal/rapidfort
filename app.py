import os
from flask import Flask, request, jsonify, render_template, send_from_directory
from werkzeug.utils import secure_filename
import mimetypes

app = Flask(__name__)

UPLOAD_FOLDER = "uploads"
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER

STATIC_FOLDER = "public"
app.config["STATIC_FOLDER"] = STATIC_FOLDER

os.makedirs(UPLOAD_FOLDER, exist_ok=True)


def extract_metadata(file_path):
    try:
        mime_details = mimetypes.guess_type(file_path)
        metadata = {
            "File Name": os.path.basename(file_path),
            "File Type": mime_details,
            "File Size": os.path.getsize(file_path),
        }

        if mime_details[0].startswith("image/"):
            add_image_metadata(file_path, metadata)
        elif mime_details[0] == "application/pdf":
            add_pdf_metadata(file_path, metadata)
        elif mime_details[0] == "text/plain":
            add_text_metadata(file_path, metadata)
        elif mime_details[0].startswith("audio/"):
            add_audio_metadata(file_path, metadata)
        elif mime_details[0].startswith("video/"):
            add_video_metadata(file_path, metadata)
        elif (
            mime_details[0]
            == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ):
            add_excel_metadata(file_path, metadata)
        elif (
            mime_details[0]
            == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ):
            add_word_document_metadata(file_path, metadata)

        return metadata
    except Exception as e:
        print(e)
        return {"Error": "Unable to extract metadata."}


# Define functions to add specific metadata based on MIME type


def add_image_metadata(file_path, metadata):
    from PIL import Image

    with Image.open(file_path) as img:
        metadata["Image Size"] = img.size
        metadata["Image Mode"] = img.mode
        metadata["Image Format"] = img.format


def add_pdf_metadata(file_path, metadata):
    from PyPDF2 import PdfReader  # Import PdfReader

    with open(file_path, "rb") as pdf_file:
        pdf_reader = PdfReader(pdf_file)
        metadata["PDF Pages"] = len(pdf_reader.pages)


def add_text_metadata(file_path, metadata):
    with open(file_path, "r") as text_file:
        metadata["Text Content"] = text_file.read()


# Add functions for audio, video, Excel, and Word document metadata here
def add_audio_metadata(file_path, metadata):
    from mutagen.mp3 import MP3

    audio = MP3(file_path)
    metadata["Audio Length"] = audio.info.length
    metadata["Audio Bitrate"] = audio.info.bitrate


def add_video_metadata(file_path, metadata):
    from mutagen.mp4 import MP4

    video = MP4(file_path)
    metadata["Video Length"] = video.info.length
    metadata["Video Bitrate"] = video.info.bitrate


def add_excel_metadata(file_path, metadata):
    import openpyxl

    wb = openpyxl.load_workbook(file_path)
    metadata["Excel Sheets"] = wb.sheetnames


def add_word_document_metadata(file_path, metadata):
    import docx

    doc = docx.Document(file_path)
    metadata["Word Document Paragraphs"] = len(doc.paragraphs)


@app.route("/", methods=["GET"])
def home():
    return render_template("index.html")


@app.route("/upload", methods=["POST"])
def upload_file():
    if "file" not in request.files:
        return jsonify({"Error": "No file part"})

    uploaded_file = request.files["file"]

    if uploaded_file.filename == "":
        return jsonify({"Error": "No selected file"})

    if uploaded_file:
        filename = secure_filename(uploaded_file.filename)
        file_path = os.path.join(app.config["UPLOAD_FOLDER"], filename)
        uploaded_file.save(file_path)

        metadata = extract_metadata(file_path)

        os.remove(file_path)

        return jsonify(metadata)


@app.route("/public/<path:filename>")
def serve_static(filename):
    return send_from_directory(app.config["STATIC_FOLDER"], filename)


if __name__ == "__main__":
    app.run(debug=True, port=5000, host="0.0.0.0")
