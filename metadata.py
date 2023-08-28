import os
import mimetypes


def extract_metadata(file_path):
    try:
        mime_details = mimetypes.guess_type(file_path)
        metadata = {
            "File Name": os.path.basename(file_path),
            "File Type": mime_details,
            "File Size": os.path.getsize(file_path),
        }

        if mime_details[0].startswith("image/"):
            add_image_metadata(file_path, metadata)
        elif mime_details[0] == "application/pdf":
            add_pdf_metadata(file_path, metadata)
        elif mime_details[0] == "text/plain":
            add_text_metadata(file_path, metadata)
        elif mime_details[0].startswith("audio/"):
            add_audio_metadata(file_path, metadata)
        elif mime_details[0].startswith("video/"):
            add_video_metadata(file_path, metadata)
        elif (
            mime_details[0]
            == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ):
            add_excel_metadata(file_path, metadata)
        elif (
            mime_details[0]
            == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ):
            add_word_document_metadata(file_path, metadata)

        return metadata
    except Exception as e:
        print(e)
        return {"Error": "Unable to extract metadata."}


def add_image_metadata(file_path, metadata):
    from PIL import Image

    with Image.open(file_path) as img:
        metadata["Image Size"] = img.size
        metadata["Image Mode"] = img.mode
        metadata["Image Format"] = img.format


def add_pdf_metadata(file_path, metadata):
    from PyPDF2 import PdfReader  # Import PdfReader

    with open(file_path, "rb") as pdf_file:
        pdf_reader = PdfReader(pdf_file)
        metadata["PDF Pages"] = len(pdf_reader.pages)


def add_text_metadata(file_path, metadata):
    with open(file_path, "r") as text_file:
        metadata["Text Content"] = text_file.read()


# Add functions for audio, video, Excel, and Word document metadata here
def add_audio_metadata(file_path, metadata):
    from mutagen.mp3 import MP3

    audio = MP3(file_path)
    metadata["Audio Length"] = audio.info.length
    metadata["Audio Bitrate"] = audio.info.bitrate


def add_video_metadata(file_path, metadata):
    from mutagen.mp4 import MP4

    video = MP4(file_path)
    metadata["Video Length"] = video.info.length
    metadata["Video Bitrate"] = video.info.bitrate


def add_excel_metadata(file_path, metadata):
    import openpyxl

    wb = openpyxl.load_workbook(file_path)
    metadata["Excel Sheets"] = wb.sheetnames


def add_word_document_metadata(file_path, metadata):
    import docx

    doc = docx.Document(file_path)
    metadata["Word Document Paragraphs"] = len(doc.paragraphs)
